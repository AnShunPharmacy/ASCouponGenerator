import openpyxl
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import Font
import json

page_break_value = json.load(open('./assets/documents/setting.json','r'))['page_break']


wb = openpyxl.Workbook()

sheet = wb.active 

sheet.column_dimensions['A'].width = 20
sheet.column_dimensions['B'].width = 7
sheet.column_dimensions['C'].width = 15
sheet.column_dimensions['D'].width = 15
sheet.column_dimensions['E'].width = 15
sheet.column_dimensions['F'].width = 7

title_element = {
    1: '現金卷編號',
    2: '日期',
    3: '姓名',
    4: '電話',
    5: '消費金額',
    6: '回收'
}

thin_border = Border(left=Side(style='thin'), 
                     right=Side(style='thin'), 
                     top=Side(style='thin'), 
                     bottom=Side(style='thin'))
fontStyle = Font(size = "16")

def generate_title(row):
    title_counting = 1
    for title in range(1 ,7) :
        sheet.cell(row = row, column = title, value = title_element[title_counting]).font = fontStyle
        title_counting +=1

def generate_table(amount: int, serial_start_number: int, serial_prefix: str, path:str):
    generate_title(1)
    i_delay = 1
    row_count = 1
    for i in range(serial_start_number, serial_start_number + amount):
        serial_number = f'NO. {serial_prefix}{str(i).zfill(4)}'
        # print(row_count)
        if row_count % page_break_value == 0:
            row_count = 1
            i_delay +=1
            generate_title(i - 1 + i_delay)
        row_count += 1
        sheet.row_dimensions[i + i_delay].height = 30
        sheet.cell(row = i + i_delay, column = 1, value = serial_number).font = fontStyle

    for sheet_roll in range(1 ,amount + i_delay + 1) :
        sheet.cell(row=sheet_roll, column=1).border = thin_border
        for sheet_colum in range(1 ,7) :
            sheet.cell(row=sheet_roll, column=sheet_colum).border = thin_border
    wb.save(path)  

if __name__ == '__main__':
    generate_table(30, 1, '1123', 'output.xlsx')






